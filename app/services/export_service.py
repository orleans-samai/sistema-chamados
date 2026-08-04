"""Exportação de chamados para Excel (openpyxl) e PDF (reportlab)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADERS = [
    "Número", "Título", "Categoria", "Prioridade", "Status",
    "Solicitante", "Analista", "Abertura", "Encerramento",
]


def _rows(tickets):
    rows = []
    for t in tickets:
        rows.append([
            t.number or "",
            t.title,
            t.category.name if t.category else "",
            t.priority,
            t.status,
            t.requester.name if t.requester else "",
            t.assignee.name if t.assignee else "—",
            t.created_at.strftime("%d/%m/%Y %H:%M") if t.created_at else "",
            t.closed_at.strftime("%d/%m/%Y %H:%M") if t.closed_at else "—",
        ])
    return rows


def tickets_to_excel(tickets):
    """Gera uma planilha .xlsx em memória com a lista de chamados."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Chamados"

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header_font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in _rows(tickets):
        ws.append(row)

    widths = [14, 40, 18, 12, 18, 22, 22, 17, 17]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def tickets_to_pdf(tickets):
    """Gera um relatório PDF em memória com a lista de chamados."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1 * cm,
        rightMargin=1 * cm,
        topMargin=1 * cm,
        bottomMargin=1 * cm,
        title="Relatório de Chamados",
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 7.5
    cell_style.leading = 9

    elements = [
        Paragraph("Relatório de Chamados — Service Desk", styles["Title"]),
        Paragraph(
            f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} — {len(tickets)} chamado(s)",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]

    data = [HEADERS] + [
        [Paragraph(str(value), cell_style) for value in row] for row in _rows(tickets)
    ]
    table = Table(
        data,
        repeatRows=1,
        colWidths=[2.3 * cm, 6.5 * cm, 2.8 * cm, 2 * cm, 2.6 * cm, 3.2 * cm, 3.2 * cm, 2.5 * cm, 2.5 * cm],
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer
